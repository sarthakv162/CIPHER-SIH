"""Phase 4 verify: every artefact fixture renders to its formats and the binaries open."""

from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path

import pytest

from rupantar.core.artefacts import ARTEFACT_MODELS
from rupantar.render.base import FORMATS, render

_TYPES = sorted(ARTEFACT_MODELS)


def _load(artefact_type: str, artefacts_dir: Path):  # type: ignore[no-untyped-def]
    model = ARTEFACT_MODELS[artefact_type]
    return model.model_validate_json((artefacts_dir / f"{artefact_type}.json").read_text())


@pytest.mark.parametrize("artefact_type", _TYPES)
def test_all_declared_formats_are_written(
    artefact_type: str, artefacts_dir: Path, tmp_path: Path
) -> None:
    """render() produces one non-empty file per format in FORMATS."""
    artefact = _load(artefact_type, artefacts_dir)
    paths = render(artefact, tmp_path)
    for path in paths:
        assert path.is_file() and path.stat().st_size > 0
    suffixes = {p.suffix.lstrip(".") for p in paths}
    if artefact_type == "video_package":
        assert {"md", "srt"} <= suffixes
        assert (tmp_path / "storyboard.json").is_file()
    else:
        assert suffixes == set(FORMATS[artefact_type])


def test_advisory_binaries_open(artefacts_dir: Path, tmp_path: Path) -> None:
    """The advisory docx and pdf open through their libraries and carry the advisory content."""
    import docx
    from pypdf import PdfReader

    advisory = _load("advisory", artefacts_dir)
    render(advisory, tmp_path)

    document = docx.Document(str(tmp_path / "advisory.docx"))
    text = "\n".join(p.text for p in document.paragraphs)
    assert advisory.indicators[0].value in text
    assert advisory.recommended_actions[0].action in text

    pdf_bytes = (tmp_path / "advisory.pdf").read_bytes()
    assert pdf_bytes.startswith(b"%PDF") and pdf_bytes.rstrip().endswith(b"%%EOF")
    reader = PdfReader(str(tmp_path / "advisory.pdf"))
    assert len(reader.pages) >= 1
    assert advisory.advisory_id in "".join(page.extract_text() for page in reader.pages)


def test_executive_summary_docx_opens(artefacts_dir: Path, tmp_path: Path) -> None:
    """The executive summary docx opens through python-docx."""
    import docx

    artefact = _load("executive_summary", artefacts_dir)
    render(artefact, tmp_path)
    document = docx.Document(str(tmp_path / "executive_summary.docx"))
    assert artefact.title in "\n".join(p.text for p in document.paragraphs)


def test_presentation_pptx_slides_and_notes(artefacts_dir: Path, tmp_path: Path) -> None:
    """The pptx opens: title slide + one per Slide, each content slide titled, bulleted, noted."""
    import pptx

    from rupantar.render.pptx_render import _body_placeholder

    deck = _load("presentation", artefacts_dir)
    render(deck, tmp_path)
    opened = pptx.Presentation(str(tmp_path / "presentation.pptx"))
    slides = list(opened.slides)
    assert len(slides) == len(deck.slides) + 1
    for rendered, spec in zip(slides[1:], deck.slides, strict=True):
        assert rendered.shapes.title is not None and rendered.shapes.title.text == spec.title
        body = _body_placeholder(rendered)
        assert body is not None
        assert any(p.text.strip() for p in body.text_frame.paragraphs)
        assert rendered.notes_slide.notes_text_frame.text.strip() == spec.speaker_notes.strip()


def test_infographic_svg_is_xml(artefacts_dir: Path, tmp_path: Path) -> None:
    """The infographic SVG parses as XML with an <svg> root."""
    spec = _load("infographic_spec", artefacts_dir)
    render(spec, tmp_path)
    root = ET.fromstring((tmp_path / "infographic_spec.svg").read_text(encoding="utf-8"))
    assert root.tag.endswith("svg")


def test_video_package_srt_has_cues(artefacts_dir: Path, tmp_path: Path) -> None:
    """The video package SRT has one cue per scene."""
    video = _load("video_package", artefacts_dir)
    render(video, tmp_path)
    srt = (tmp_path / "video_package.srt").read_text(encoding="utf-8")
    assert "-->" in srt
    assert srt.count("-->") == len(video.scenes)
