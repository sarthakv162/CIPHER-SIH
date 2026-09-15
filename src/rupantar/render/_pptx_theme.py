"""Colour/shape decoration layer for the pptx renderer: title-run tinting, an accent-bar
strip, and the quote/closing slide treatments -- all driven by the same `Theme` the
video/SVG/PDF renderers use. Every function degrades silently on its own (the caller wraps
each call so one broken decoration never blocks the rest of the slide or the deck).
"""

from __future__ import annotations

from typing import Any

from rupantar.render.theme import Theme

RGB = tuple[int, int, int]


def colour_title(slide: Any, rgb: RGB) -> None:
    """Tint every run of the slide's title placeholder, when one is present."""
    from pptx.dml.color import RGBColor

    title = slide.shapes.title
    if title is None:
        return
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(*rgb)


def accent_bar(deck: Any, slide: Any, theme: Theme) -> None:
    """A thin accent-coloured strip along the top edge of a slide."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu

    height = Emu(91440)  # 0.1 inch, independent of the deck's own slide size
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*theme.rgb("accent"))
    bar.line.fill.background()
    bar.shadow.inherit = False


def quote_treatment(deck: Any, slide: Any, body: Any, theme: Theme) -> None:
    """Centre/italicise/enlarge a quote slide's body text and add a large accent quote mark."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Pt

    if body is not None and body.has_text_frame:
        for paragraph in body.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.italic = True
                run.font.size = Pt(theme.px("heading"))
                run.font.color.rgb = RGBColor(*theme.rgb("primary"))
    mark = slide.shapes.add_textbox(
        Emu(int(deck.slide_width * 0.04)),
        Emu(int(deck.slide_height * 0.16)),
        Emu(int(deck.slide_width * 0.15)),
        Emu(int(deck.slide_height * 0.2)),
    )
    run = mark.text_frame.paragraphs[0].add_run()
    run.text = "“"
    run.font.size = Pt(96)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*theme.rgb("accent"))


def closing_treatment(slide: Any, body: Any, theme: Theme) -> None:
    """Fill a closing slide's background with the primary colour and invert its text."""
    from pptx.dml.color import RGBColor

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*theme.rgb("primary"))
    colour_title(slide, theme.rgb("text"))
    if body is not None and body.has_text_frame:
        for paragraph in body.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(*theme.rgb("text"))
