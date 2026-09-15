"""PowerPoint rendering for a Presentation artefact: title slide plus one slide per Slide.

Loads the selected template's `.potx` via python-pptx and maps each `Slide.layout` enum value
(plus the deck's own opening title slide) to a layout **index** declared in
`configs/templates/templates.yaml` -- never a guessed position. A missing template, a missing
or unopenable `.potx` file, or a missing/out-of-range layout index degrades to python-pptx's
built-in default deck and the nearest available layout, recording a warning on `context`
(see `render/context.py`) instead of raising. See `docs/TEMPLATES.md` for what a template needs.

On top of the template layer, `render/_pptx_theme.py` adds a colour/shape design layer --
tinted titles, an accent-bar strip, quote/closing slide treatments, and (when `context`
carries an `infographic_spec` with genuine numbers) an extra data-chart slide reusing
`render/charts.py`. Every decoration degrades independently and never raises.

After saving, `_repair_ooxml` patches `ppt/presentation.xml`: python-pptx never writes the
`<p:notesMasterIdLst>` element even when it adds a notes master, which Keynote's strict
importer rejects outright ("file format is invalid"). It also drops the deprecated
`<p:sldSz type=...>` attribute. PowerPoint / LibreOffice tolerate both; Keynote does not.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import Any

from rupantar.render import _pptx_theme as decor
from rupantar.render._theme_fallback import safe_theme
from rupantar.render.context import RenderContext
from rupantar.render.template_registry import DEFAULT_TEMPLATE, get_template, template_file
from rupantar.render.theme import Theme

# Used only when no template is engaged, or the engaged one has no usable pptx binding --
# indices into python-pptx's own built-in default deck. Content slides never use layout 0
# (the Title Slide layout) -- the deck already has one, added by _title_slide.
_FALLBACK_LAYOUTS = {
    "title_slide": 0,
    "title": 2,
    "bullets": 1,
    "two_column": 3,
    "quote": 1,
    "closing": 2,
}
_REQUIRED_LAYOUTS = ("title_slide", "title", "bullets", "two_column", "quote", "closing")
_NOTES_MASTER_RID = re.compile(
    r'<Relationship\b(?=[^>]*\bType="[^"]*/notesMaster")[^>]*\bId="(rId\d+)"'
)
_SLD_SZ_TYPE = re.compile(r'(<p:sldSz\b[^>]*?)\s+type="[^"]*"([^>]*/>)')


def render_pptx(artefact: Any, path: Path, *, context: RenderContext | None = None) -> list[Path]:
    """Render the deck to a .pptx file (+ a sibling chart PNG, when one is drawn)."""
    from pptx import Presentation as PptxPresentation

    theme = safe_theme(context)
    template_path, template_layouts = _resolve_template(context)
    deck, layout_map = _open_deck(PptxPresentation, template_path, template_layouts, context)
    layouts = _resolve_layouts(layout_map, len(deck.slide_layouts), context)
    _title_slide(deck, artefact.title, artefact.deck_summary, layouts, theme, context)

    slides = list(artefact.slides)
    closing_last = bool(slides) and slides[-1].layout.value == "closing"
    body_slides = slides[:-1] if closing_last else slides
    for slide_spec in body_slides:
        _content_slide(deck, slide_spec, layouts, theme, context)
    chart_path = _maybe_chart_slide(deck, theme, layouts, context, path)
    if closing_last:
        _content_slide(deck, slides[-1], layouts, theme, context)

    deck.save(str(path))
    _repair_ooxml(path)
    return [path, chart_path] if chart_path is not None else [path]


def _resolve_template(context: RenderContext | None) -> tuple[Path | None, dict[str, int]]:
    """The `.potx` file and its declared layout map, or (None, {}) to use the built-in deck."""
    template_id = context.template_id if context else DEFAULT_TEMPLATE
    configs_dir = context.configs_dir if context else None
    spec = get_template(template_id, configs_dir=configs_dir)
    if spec is None:
        if template_id != DEFAULT_TEMPLATE:
            _warn(context, f"template {template_id!r} is not defined; using the built-in deck")
        return None, {}
    if spec.pptx is None:
        _warn(context, f"template {template_id!r} has no pptx binding; using the built-in deck")
        return None, {}
    path = template_file(configs_dir, spec.pptx.file)
    if not path.is_file():
        _warn(context, f"template {template_id!r}'s file {spec.pptx.file!r} is missing")
        return None, {}
    return path, dict(spec.pptx.layouts)


def _open_deck(
    pptx_presentation: Any,
    template_path: Path | None,
    template_layouts: dict[str, int],
    context: RenderContext | None,
) -> tuple[Any, dict[str, int]]:
    """Open `template_path` if given, else the built-in deck; return it with its layout map."""
    if template_path is not None:
        try:
            return pptx_presentation(str(template_path)), template_layouts
        except Exception as exc:  # a corrupt or unreadable .potx must never crash the job
            _warn(context, f"template file {template_path.name!r} could not be opened ({exc})")
    return pptx_presentation(), dict(_FALLBACK_LAYOUTS)


def _resolve_layouts(
    layout_map: dict[str, int], layout_count: int, context: RenderContext | None
) -> dict[str, int]:
    """Every required layout key mapped to a valid index, falling back to the nearest one."""
    resolved: dict[str, int] = {}
    for key in _REQUIRED_LAYOUTS:
        index = layout_map.get(key)
        if index is not None and layout_count and 0 <= index < layout_count:
            resolved[key] = index
            continue
        fallback = min(_FALLBACK_LAYOUTS[key], max(0, layout_count - 1))
        _warn(context, f"template layout {key!r} missing or out of range; using index {fallback}")
        resolved[key] = fallback
    return resolved


def _warn(context: RenderContext | None, message: str) -> None:
    """Record a degradation warning on the shared context, if one was given."""
    if context is not None:
        context.warnings.append(message)


def _decorate_slide(
    deck: Any, slide: Any, layout_key: str, body: Any, theme: Theme, context: RenderContext | None
) -> None:
    """Apply the colour/shape design layer to one slide; a failure skips it, never the render."""
    try:
        decor.accent_bar(deck, slide, theme)
        decor.colour_title(slide, theme.rgb("primary"))
        if layout_key == "quote":
            decor.quote_treatment(deck, slide, body, theme)
        if layout_key == "closing":
            decor.closing_treatment(slide, body, theme)
    except Exception as exc:  # a decoration slip must never block the deck
        _warn(context, f"slide decoration skipped ({exc})")


def _title_slide(
    deck: Any, title: str, subtitle: str, layouts: dict[str, int], theme: Theme, context: Any
) -> None:
    """Add the opening title slide."""
    slide = deck.slides.add_slide(deck.slide_layouts[layouts["title_slide"]])
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    body = _body_placeholder(slide)
    if body is not None:
        body.text = subtitle
    _decorate_slide(deck, slide, "title_slide", body, theme, context)


def _content_slide(
    deck: Any, slide_spec: Any, layouts: dict[str, int], theme: Theme, context: Any
) -> None:
    """Add one content slide with bullets, speaker notes, and the colour/shape design layer."""
    layout_index = layouts[slide_spec.layout.value]
    slide = deck.slides.add_slide(deck.slide_layouts[layout_index])
    if slide.shapes.title is not None:
        slide.shapes.title.text = slide_spec.title
    body = _body_placeholder(slide)
    if body is not None and slide_spec.bullets:
        frame = body.text_frame
        frame.text = str(slide_spec.bullets[0])
        for bullet in slide_spec.bullets[1:]:
            frame.add_paragraph().text = str(bullet)
    slide.notes_slide.notes_text_frame.text = slide_spec.speaker_notes
    _decorate_slide(deck, slide, slide_spec.layout.value, body, theme, context)


def _maybe_chart_slide(
    deck: Any, theme: Theme, layouts: dict[str, int], context: RenderContext | None, path: Path
) -> Path | None:
    """Extract+render a data chart and insert its slide, or return None; never raises."""
    if context is None or context.infographic_spec is None:
        return None
    try:
        from rupantar.render import charts

        data = charts.extract_chart_data(context.infographic_spec)
        if data is None:
            return None
        chart_path = path.with_name(f"{path.stem}_chart.png")
        charts.render_chart(chart_path, data, theme)
        slide = deck.slides.add_slide(deck.slide_layouts[layouts["bullets"]])
        if slide.shapes.title is not None:
            slide.shapes.title.text = context.infographic_spec.headline or "Data"
        _place_chart_picture(deck, slide, chart_path)
        _decorate_slide(deck, slide, "bullets", None, theme, context)
        return chart_path
    except Exception as exc:
        _warn(context, f"chart slide could not be built ({exc})")
        return None


def _place_chart_picture(deck: Any, slide: Any, chart_path: Path) -> None:
    """Size and centre the chart PNG (1280x720) under the title, within the slide bounds."""
    from pptx.util import Emu

    margin = Emu(457200)
    top_offset = Emu(1500000)
    avail_w = deck.slide_width - 2 * margin
    avail_h = deck.slide_height - top_offset - margin
    target_w, target_h = avail_w, int(avail_w * 720 / 1280)
    if target_h > avail_h:
        target_h, target_w = avail_h, int(avail_h * 1280 / 720)
    left = (deck.slide_width - target_w) // 2
    slide.shapes.add_picture(str(chart_path), left, top_offset, width=target_w, height=target_h)


def _body_placeholder(slide: Any) -> Any:
    """Return the first non-title placeholder on a slide, or None."""
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.idx != 0:
            return placeholder
    return None


def _repair_ooxml(path: Path) -> None:
    """Rewrite the package with `ppt/presentation.xml` made acceptable to strict importers."""
    with zipfile.ZipFile(path) as archive:
        parts = [(info, archive.read(info.filename)) for info in archive.infolist()]
    rels = next(d for i, d in parts if i.filename == "ppt/_rels/presentation.xml.rels")
    match = _NOTES_MASTER_RID.search(rels.decode("utf-8"))
    notes_rid = match.group(1) if match else None

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in parts:
            if info.filename == "ppt/presentation.xml":
                data = _patch_presentation(data.decode("utf-8"), notes_rid).encode("utf-8")
            out.writestr(info, data)
    path.write_bytes(buffer.getvalue())


def _patch_presentation(xml: str, notes_rid: str | None) -> str:
    """Add `<p:notesMasterIdLst>` after the slide-master list and drop `<p:sldSz type=...>`."""
    xml = _SLD_SZ_TYPE.sub(r"\1\2", xml)
    if notes_rid and "notesMasterIdLst" not in xml:
        xml = xml.replace(
            "</p:sldMasterIdLst>",
            "</p:sldMasterIdLst>"
            f'<p:notesMasterIdLst><p:notesMasterId r:id="{notes_rid}"/></p:notesMasterIdLst>',
            1,
        )
    return xml
