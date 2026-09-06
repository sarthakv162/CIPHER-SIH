"""General-format readers: each returns (list[dict], warnings) and never raises on bad content."""

from __future__ import annotations

import csv
import json
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

import yaml

Record = dict[str, Any]


def _coerce_records(data: Any, name: str, fmt: str) -> tuple[list[Record], list[str]]:
    """Normalise a parsed document into a list of dict rows with a warning for anything dropped."""
    if isinstance(data, list):
        rows = [item for item in data if isinstance(item, dict)]
        if len(rows) != len(data):
            return rows, [
                f"{name}: {len(data) - len(rows)} non-object item(s) in {fmt} array skipped"
            ]
        return rows, []
    if isinstance(data, dict):
        return [data], []
    return [], [f"{name}: top-level {fmt} is not an object or an array of objects"]


def _read_delimited(
    path: Path, opts: dict[str, Any], fmt: str, delimiter: str
) -> tuple[list[Record], list[str]]:
    """Read csv/tsv with BOM tolerance, padding short rows and keeping long-row extras."""
    warnings: list[str] = []
    rows: list[Record] = []
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            header: list[str] | None = None
            for lineno, raw in enumerate(reader, start=1):
                if not raw or all(cell == "" for cell in raw):
                    continue
                if header is None:
                    header = raw
                    continue
                rows.append(_delimited_row(raw, header, path.name, lineno, warnings))
    except (OSError, csv.Error, UnicodeDecodeError) as exc:
        return [], [f"{path.name}: could not parse as {fmt}: {exc}"]
    return rows, warnings


def _delimited_row(
    raw: list[str], header: list[str], name: str, lineno: int, warnings: list[str]
) -> Record:
    """Build one row dict, padding a short row and stashing a long row's extras under _extra_N."""
    if len(raw) < len(header):
        warnings.append(
            f"{name}: row {lineno} has {len(raw)} fields, expected {len(header)}; padded"
        )
        raw = raw + [""] * (len(header) - len(raw))
    row: Record = {col: raw[idx] for idx, col in enumerate(header)}
    if len(raw) > len(header):
        warnings.append(
            f"{name}: row {lineno} has {len(raw)} fields, over the {len(header)}-col header;"
            " extras kept as _extra_N"
        )
        for offset, value in enumerate(raw[len(header) :], start=1):
            row[f"_extra_{offset}"] = value
    return row


def read_csv(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read a comma-separated file into rows."""
    return _read_delimited(path, opts, "csv", ",")


def read_tsv(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read a tab-separated file into rows."""
    return _read_delimited(path, opts, "tsv", "\t")


def read_json(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read a JSON array of objects or a single object."""
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError, UnicodeDecodeError) as exc:
        return [], [f"{path.name}: could not parse as json: {exc}"]
    return _coerce_records(data, path.name, "json")


def read_jsonl(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read one JSON object per line, skipping malformed lines with a warning."""
    warnings: list[str] = []
    rows: list[Record] = []
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeDecodeError) as exc:
        return [], [f"{path.name}: could not parse as jsonl: {exc}"]
    for lineno, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue
        try:
            obj = json.loads(line)
        except json.JSONDecodeError as exc:
            warnings.append(f"{path.name}: line {lineno} is not valid JSON ({exc}); skipped")
            continue
        if isinstance(obj, dict):
            rows.append(obj)
        else:
            warnings.append(f"{path.name}: line {lineno} is not a JSON object; skipped")
    return rows, warnings


def read_yaml(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read a YAML list of mappings or a single mapping."""
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, yaml.YAMLError) as exc:
        return [], [f"{path.name}: could not parse as yaml: {exc}"]
    return _coerce_records(data, path.name, "yaml")


def read_xml(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read <rows><row><col name="k">v</col></row></rows>, tolerating a bare root or attributes."""
    warnings: list[str] = []
    try:
        root = ET.parse(path).getroot()
    except (OSError, ET.ParseError) as exc:
        return [], [f"{path.name}: could not parse as xml: {exc}"]
    row_elems = root.findall("row")
    if not row_elems and root.tag == "row":
        warnings.append(f"{path.name}: no <rows> wrapper; treated the root <row> as a single row")
        row_elems = [root]
    rows: list[Record] = []
    for elem in row_elems:
        row: Record = {}
        children = list(elem)
        for child in children:
            row[child.get("name", child.tag)] = child.text or ""
        if not children and elem.attrib:
            warnings.append(f"{path.name}: <row> has no <col> children; used its attributes")
            row.update(elem.attrib)
        rows.append(row)
    return rows, warnings


def read_xlsx(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read the first sheet of a workbook, first row as the header."""
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        header: list[str] | None = None
        rows: list[Record] = []
        for raw in sheet.iter_rows(values_only=True):
            if raw is None or all(cell is None for cell in raw):
                continue
            if header is None:
                header = [
                    str(cell) if cell is not None else f"col_{idx}" for idx, cell in enumerate(raw)
                ]
                continue
            rows.append(
                {
                    col: "" if idx >= len(raw) or raw[idx] is None else str(raw[idx])
                    for idx, col in enumerate(header)
                }
            )
        workbook.close()
    except Exception as exc:  # openpyxl raises many unrelated types on bad input
        return [], [f"{path.name}: could not parse as xlsx: {exc}"]
    return rows, []


def read_parquet(path: Path, opts: dict[str, Any]) -> tuple[list[Record], list[str]]:
    """Read a Parquet file, coercing every value to a string for round-trip fidelity."""
    try:
        import pyarrow.parquet as pq

        table = pq.read_table(path)
        columns = table.column_names
        rows = [
            {col: "" if item.get(col) is None else str(item[col]) for col in columns}
            for item in table.to_pylist()
        ]
    except Exception as exc:  # pyarrow raises many unrelated types on bad input
        return [], [f"{path.name}: could not parse as parquet: {exc}"]
    return rows, []
