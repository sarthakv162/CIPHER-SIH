#!/usr/bin/env python3
"""Generate placeholder .potx/.dotx/thumbnail files for configs/templates/templates.yaml.

Run once, offline, to scaffold the three starter templates (ntro-formal, executive,
technical) with structurally correct layouts and named styles. These are deliberately
plain placeholders -- real visual design is expected to replace them; see docs/TEMPLATES.md
for exactly which layout indices and style names a replacement file must keep.

Usage: .venv/bin/python scripts/generate_placeholder_templates.py
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

_REPO = Path(__file__).resolve().parents[1]
_TEMPLATES_DIR = _REPO / "configs" / "templates"

# dk2 drives many masters' title-text colour; accent1 drives shape/line accents. Recolouring
# both is the standard way a real theme differentiates -- not a renderer-side hack.
_THEME_TAGS = ("dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6")


def _recolor_theme(potx_path: Path, accent_hex: str) -> None:
    """Rewrite ppt/theme/theme1.xml's dk2/accentN colours in a saved .potx, in place."""
    with zipfile.ZipFile(potx_path) as archive:
        parts = [(info, archive.read(info.filename)) for info in archive.infolist()]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in parts:
            if info.filename == "ppt/theme/theme1.xml":
                xml = data.decode("utf-8")
                for tag in _THEME_TAGS:
                    xml = re.sub(
                        rf'(<a:{tag}><a:srgbClr val=")[0-9A-Fa-f]{{6}}("/></a:{tag}>)',
                        rf"\g<1>{accent_hex}\g<2>",
                        xml,
                    )
                data = xml.encode("utf-8")
            out.writestr(info, data)
    potx_path.write_bytes(buffer.getvalue())

# (id, label, accent RGB) -- the built-in python-pptx/python-docx default templates already
# have every slide layout and paragraph style docs/TEMPLATES.md requires except the two
# advisory-specific docx styles, which are added below. Only the accent colour and the two
# custom styles vary between the three starter templates.
_TEMPLATES = [
    ("ntro-formal", "NTRO Formal", (15, 55, 90)),
    ("executive", "Executive", (30, 90, 70)),
    ("technical", "Technical", (70, 40, 110)),
]


def _make_pptx(template_id: str, accent: tuple[int, int, int]) -> None:
    """A .potx built on python-pptx's default deck, which already has every required layout.

    Deliberately zero slides: render_pptx opens this file and *appends* the deck's title
    slide plus one slide per Slide spec, so any slide already present here would leak into
    every rendered .pptx as unwanted leading content. The theme's dk2/accent colours are
    recoloured after saving so the three placeholders are not byte-identical decks --
    the accent is what a real template swap should visibly change.
    """
    from pptx import Presentation

    deck = Presentation()
    assert len(deck.slides) == 0, "a placeholder template must start with no slides"
    out = _TEMPLATES_DIR / f"{template_id}.potx"
    deck.save(str(out))
    _recolor_theme(out, "%02X%02X%02X" % accent)
    print(f"wrote {out} ({len(deck.slide_layouts)} layouts, {len(deck.slides)} slides)")


def _add_style(
    document: object,
    name: str,
    *,
    bold: bool,
    italic: bool,
    size_pt: float,
    rgb: tuple[int, int, int],
) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt, RGBColor

    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)  # type: ignore[attr-defined]
    style.base_style = document.styles["Normal"]  # type: ignore[attr-defined]
    style.font.bold = bold
    style.font.italic = italic
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(*rgb)


def _make_docx(template_id: str, accent: tuple[int, int, int]) -> None:
    """A .dotx with two custom styles added: Severity Banner and Classification Footer.

    Deliberately zero body paragraphs: render_docx opens this file and *appends* every
    paragraph it writes, so any paragraph already present here would leak into every
    rendered .docx as unwanted leading content.
    """
    from docx import Document

    document = Document()
    assert len(document.paragraphs) == 0, "a placeholder template must start with no paragraphs"
    _add_style(document, "Severity Banner", bold=True, italic=False, size_pt=13, rgb=accent)
    grey = (110, 110, 110)
    _add_style(document, "Classification Footer", bold=False, italic=True, size_pt=8, rgb=grey)

    out = _TEMPLATES_DIR / f"{template_id}.dotx"
    document.save(str(out))
    added = [s.name for s in document.styles if "Banner" in s.name or "Footer" in s.name]
    print(f"wrote {out} (styles: {added}, {len(document.paragraphs)} paragraphs)")


def _make_thumbnail(template_id: str, label: str, accent: tuple[int, int, int]) -> None:
    """A small solid-colour PNG standing in for a real preview render."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (320, 200), accent)
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 160, 320, 200), fill=(255, 255, 255))
    draw.text((16, 170), label, fill=(20, 20, 20))
    draw.text((16, 16), "PLACEHOLDER", fill=(255, 255, 255))
    out = _TEMPLATES_DIR / f"{template_id}.png"
    image.save(out, format="PNG")
    print(f"wrote {out}")


def main() -> None:
    _TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    for template_id, label, accent in _TEMPLATES:
        _make_pptx(template_id, accent)
        _make_docx(template_id, accent)
        _make_thumbnail(template_id, label, accent)


if __name__ == "__main__":
    main()
